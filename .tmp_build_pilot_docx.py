from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\Programming\Projects\ArtistCRM_PartyCRM\PartyCRM")
SOURCE = ROOT / "docs" / "PARTYCRM_USER_PILOT_CHECKLIST.md"
OUTPUT = ROOT / "docs" / "PartyCRM_Чеклист_для_компании.docx"

NAVY = RGBColor(11, 62, 93)
BLUE = RGBColor(21, 101, 154)
MUTED = RGBColor(91, 103, 112)
LIGHT = "EAF4FA"
PALE = "F5F8FA"
WHITE = RGBColor(255, 255, 255)


def set_font(run, size=11, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=140, bottom=90, end=140):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Страница ")
    set_font(run, 9, MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for style_name, size, color, before, after in (
    ("Heading 1", 17, NAVY, 16, 8),
    ("Heading 2", 14, BLUE, 13, 7),
    ("Heading 3", 12, NAVY, 10, 5),
):
    style = styles[style_name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

check_style = styles.add_style("Pilot Checklist Item", WD_STYLE_TYPE.PARAGRAPH)
check_style.base_style = normal
check_style.font.name = "Calibri"
check_style.font.size = Pt(11)
check_style.paragraph_format.left_indent = Inches(0.28)
check_style.paragraph_format.first_line_indent = Inches(-0.28)
check_style.paragraph_format.space_before = Pt(1.5)
check_style.paragraph_format.space_after = Pt(5)
check_style.paragraph_format.line_spacing = 1.15

response_style = styles.add_style("Pilot Response", WD_STYLE_TYPE.PARAGRAPH)
response_style.base_style = normal
response_style.font.name = "Calibri"
response_style.font.size = Pt(10.5)
response_style.font.color.rgb = MUTED
response_style.paragraph_format.space_before = Pt(2)
response_style.paragraph_format.space_after = Pt(7)

header = section.header
hp = header.paragraphs[0]
hp.text = "PARTYCRM  |  ПИЛОТНАЯ ПРОГРАММА"
set_font(hp.runs[0], 9, MUTED, bold=True)

footer = section.footer
fp = footer.paragraphs[0]
add_page_number(fp)

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(14)
p.paragraph_format.space_after = Pt(4)
r = p.add_run("ЧЕКЛИСТ ДЛЯ КОМПАНИИ")
set_font(r, 10, BLUE, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Проверка PartyCRM в ежедневной работе")
set_font(r, 27, NAVY, bold=True)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(16)
r = p.add_run("Простой бланк для сотрудников, администраторов и исполнителей")
set_font(r, 13, MUTED)

meta = doc.add_table(rows=2, cols=2)
meta.autofit = False
widths = (Inches(3.35), Inches(3.35))
labels = (("Компания", "Дата заполнения"), ("Имя сотрудника", "Роль"))
for row, row_labels in zip(meta.rows, labels):
    for cell, label, width in zip(row.cells, row_labels, widths):
        cell.width = width
        set_cell_margins(cell)
        shade_paragraph(cell.paragraphs[0], PALE)
        run = cell.paragraphs[0].add_run(f"{label}: __________________________________")
        set_font(run, 10.5, NAVY, bold=True)

doc.add_paragraph()

intro = doc.add_paragraph()
intro.paragraph_format.space_before = Pt(3)
intro.paragraph_format.space_after = Pt(8)
shade_paragraph(intro, LIGHT)
r = intro.add_run(
    "Это не экзамен и не техническое тестирование. Работайте как обычно и отмечайте только то, что действительно попробовали."
)
set_font(r, 11, NAVY, bold=True)

legend = doc.add_paragraph()
legend.paragraph_format.space_after = Pt(12)
for text, bold in (
    ("☒ получилось   ", True),
    ("⚠ получилось, но неудобно   ", True),
    ("☐ ещё не пробовали   ", True),
    ("— не требуется", True),
):
    r = legend.add_run(text)
    set_font(r, 10.5, NAVY if bold else MUTED, bold=bold)

lines = SOURCE.read_text(encoding="utf-8").splitlines()
started = False
for raw in lines:
    line = raw.strip()
    if line.startswith("## 1."):
        started = True
    if not started or not line:
        continue
    if line.startswith("## "):
        p = doc.add_paragraph(line[3:], style="Heading 1")
        p.paragraph_format.page_break_before = line.startswith("## 9.")
    elif line.startswith("### "):
        doc.add_paragraph(line[4:], style="Heading 2")
    elif line.startswith("- [ ] "):
        p = doc.add_paragraph(style="Pilot Checklist Item")
        r = p.add_run("☐  " + line[6:])
        set_font(r, 11, NAVY)
    elif line.startswith("- "):
        p = doc.add_paragraph(style="Pilot Checklist Item")
        r = p.add_run("•  " + line[2:])
        set_font(r, 11)
    elif line == ">":
        for _ in range(2):
            p = doc.add_paragraph("________________________________________________________________________________", style="Pilot Response")
    elif line.endswith(":") or line.startswith("Дата заполнения") or line.startswith("Роль сотрудника") or line.startswith("Устройство") or line.startswith("Как часто") or line.startswith("Оцените") or line.startswith("Стали бы"):
        p = doc.add_paragraph(style="Pilot Response")
        r = p.add_run(line + "  " + "_" * 42)
        set_font(r, 10.5, MUTED, bold=True)
    elif line[0:2] in ("1.", "2.", "3."):
        p = doc.add_paragraph(style="Pilot Response")
        r = p.add_run(line + " " + "_" * 70)
        set_font(r, 10.5, MUTED)
    else:
        p = doc.add_paragraph(style="Pilot Response")
        r = p.add_run(line)
        set_font(r, 10.5, MUTED)

doc.core_properties.title = "PartyCRM — чеклист для компании"
doc.core_properties.subject = "Пользовательская проверка PartyCRM в ежедневной работе"
doc.core_properties.author = "PartyCRM"
doc.save(OUTPUT)
print(OUTPUT)
